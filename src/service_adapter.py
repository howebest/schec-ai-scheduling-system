# -*- coding: utf-8 -*-
"""Rust-managed JSON adapter for the existing SCHEC scheduling CLI.

The adapter is intentionally dependency-light at import time so request
validation and diagnostics work even when the optional solver stack is absent.
"""
from __future__ import annotations

import argparse
import contextlib
import csv
import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import sys
import time
from pathlib import Path
from typing import Any


REQUEST_SCHEMA_VERSION = 1
RESPONSE_SCHEMA = "schec.job-response/v1"
SUPPORTED_TASKS = {
    "info", "solve_demo", "solve_full", "verify_baseline", "reschedule",
    "kpi", "validate_schedule", "preview_import", "adjust_schedule",
}
SCHEDULE_OBJECTIVE_PROFILES = {"balanced", "coverage_first", "cost_control", "fairness_first"}
EVENT_TYPES = {"insert", "rampup", "changeover", "equipment_fail", "absence"}
EVENT_REQUIRED = {
    "insert": ("line", "date", "shift"),
    "rampup": ("line", "date", "add_per_position"),
    "changeover": ("line", "date", "shift"),
    "equipment_fail": ("line", "dates"),
    "absence": ("employee", "date"),
}
TASK_FILES = {
    "info": {"input"},
    "solve_demo": {"input"},
    "solve_full": {"input"},
    "verify_baseline": {"input", "baseline"},
    "reschedule": {"input", "baseline"},
    "kpi": {"baseline"},
    "validate_schedule": {"input", "schedule"},
    "preview_import": {"upload"},
    "adjust_schedule": {"input", "schedule"},
}
IMPORT_SCHEMAS = {
    "orders": {"required": ("order_id", "line_code", "product_code", "required_units", "due_at", "priority"), "extension": ".csv"},
    "line_capacity": {"required": ("line_code", "product_code", "units_per_hour", "efficiency_pct", "available_hours", "shift_code"), "extension": ".csv"},
    "attendance_snapshot": {"required": ("employee_id", "schedule_date", "shift_code", "status", "source_type"), "extension": ".csv"},
    "position_requirements": {"required": ("工厂代码", "产线编码", "岗位编码", "岗位名称", "岗位类别", "岗位技能要求等级", "白班定编", "夜班定编", "轮班模式"), "extension": ".csv"},
    "line_run_plan": {"required": ("产线编码", "日期", "班次", "生产运行状态"), "extension": ".csv"},
    "baseline_schedule": {"required": ("工厂代码", "产线编码", "日期", "班次", "班次代码", "员工工号", "岗位编码", "岗位名称", "是否关键岗", "技能等级", "计薪工时(h)", "标准工时(h)", "当日加班工时(h)", "该班次换型等级", "实验编号"), "extension": ".csv"},
    "raw_dataset": {"extension": ".xlsx"},
    "event_json": {"extension": ".json"},
}
RAW_WORKBOOK_SHEETS = {
    "设备产线与工位台账": ("工厂代码", "产线编码", "岗位编码", "岗位名称", "岗位类别", "岗位技能要求等级", "白班定编", "夜班定编", "轮班模式"),
    "人员技能矩阵": ("员工工号", "工厂代码", "技能岗位", "技能等级", "证书有效期", "可胜任班次", "主岗标记", "累计在岗时长"),
    "排班训练主数据集": ("工厂代码", "产线编码", "日期", "班次", "生产运行状态", "换型标记", "换型等级", "岗位编码", "持证要求"),
}
HARD_CONSTRAINT_LABELS = {
    "qual_night": "岗位资格或夜班资格不满足",
    "one_per_day": "同一员工同一日期存在多班指派",
    "interval_11h": "相邻班次间隔不足 11 h",
    "rest_7d": "连续 7 日内休息日不足 1 日",
    "month_ot_36h": "月加班超过 36 h",
    "daily_hours": "日标准工时或加班超过限制",
    "coverage_eq": "岗位需求覆盖账目不平衡",
}
CSV_UNIQUE_FIELDS = {
    "orders": ("order_id",),
    "line_capacity": ("line_code", "product_code", "shift_code"),
    "attendance_snapshot": ("employee_id", "schedule_date", "shift_code"),
    "position_requirements": ("岗位编码",),
    "line_run_plan": ("产线编码", "日期", "班次"),
    "baseline_schedule": ("员工工号", "岗位编码", "日期", "班次代码"),
}
JOB_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,63}$")


class RequestError(ValueError):
    def __init__(self, code: str, message: str, details: list[str] | None = None):
        super().__init__(message)
        self.code = code
        self.details = details or []


def _response(job_id: str, status: str, exit_code: int,
              summary: Any = None, verification: Any = None,
              timings: dict[str, Any] | None = None,
              artifacts: list[dict[str, Any]] | None = None,
              error: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "schema": RESPONSE_SCHEMA,
        "job_id": job_id,
        "status": status,
        "exit_code": int(exit_code),
        "summary": summary,
        "verification": verification,
        "timings": timings or {"wall_seconds": 0.0},
        "artifacts": artifacts or [],
        "error": error,
    }


def _error_response(job_id: str, code: str, message: str,
                    details: list[str] | None = None) -> dict[str, Any]:
    return _response(
        job_id=job_id,
        status="failed",
        exit_code=2,
        error={"code": code, "message": message, "details": details or []},
    )


def _safe_job_id(value: Any) -> str:
    if isinstance(value, str) and JOB_ID_PATTERN.fullmatch(value) and ".." not in value:
        return value
    return "invalid-request"


def _resolve_package_file(package_root: Path, value: Any, field: str) -> Path:
    if not isinstance(value, str) or not value.strip():
        raise RequestError("INVALID_REQUEST", f"字段 input_files.{field} 必须为非空相对路径。")
    path_text = value.strip()
    if "\\" in path_text:
        raise RequestError("PATH_OUTSIDE_PACKAGE", f"字段 input_files.{field} 必须使用包内 POSIX 相对路径。")
    rel = Path(path_text)
    if rel.is_absolute() or any(part in ("..", "") for part in rel.parts):
        raise RequestError("PATH_OUTSIDE_PACKAGE", f"字段 input_files.{field} 超出工程包目录。")
    root = package_root.resolve(strict=True)
    candidate = (root / rel).resolve(strict=True)
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise RequestError("PATH_OUTSIDE_PACKAGE", f"字段 input_files.{field} 超出工程包目录。") from exc
    if not candidate.is_file():
        raise RequestError("INVALID_REQUEST", f"字段 input_files.{field} 必须指向文件。")
    return candidate


def _parse_date(value: Any, field: str) -> dt.date:
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise RequestError("INVALID_EVENT", f"事件字段 {field} 必须为 YYYY-MM-DD 日期。")
    try:
        return dt.date.fromisoformat(value)
    except ValueError as exc:
        raise RequestError("INVALID_EVENT", f"事件字段 {field} 不是有效日期。") from exc


def _validate_event(event: Any) -> dict[str, Any]:
    if not isinstance(event, dict):
        raise RequestError("INVALID_EVENT", "重排任务必须提供 event 对象。")
    event_type = event.get("type")
    if event_type not in EVENT_TYPES:
        raise RequestError("INVALID_EVENT", "事件类型不受支持。", ["支持 insert、rampup、changeover、equipment_fail、absence。"])
    missing = [key for key in EVENT_REQUIRED[event_type] if key not in event]
    if missing:
        raise RequestError("INVALID_EVENT", "事件缺少必填字段。", [", ".join(missing)])
    if not isinstance(event.get("event_id", "EVT"), str) or not event.get("event_id", "EVT").strip():
        raise RequestError("INVALID_EVENT", "event_id 必须为非空文本。")
    if event_type in {"insert", "rampup", "changeover", "absence"}:
        _parse_date(event["date"], "date")
    if event_type in {"insert", "changeover", "equipment_fail", "rampup"}:
        if not isinstance(event.get("line"), str) or not event["line"].strip():
            raise RequestError("INVALID_EVENT", "事件字段 line 必须为非空产线编码。")
    if event_type in {"insert", "changeover"} and event.get("shift") not in {"白班", "夜班", "常白班"}:
        raise RequestError("INVALID_EVENT", "事件字段 shift 必须为白班、夜班或常白班。")
    if event_type == "rampup":
        add = event.get("add_per_position")
        if isinstance(add, bool) or not isinstance(add, int) or add < 0:
            raise RequestError("INVALID_EVENT", "增产事件的 add_per_position 必须为非负整数。")
    if event_type == "equipment_fail":
        dates = event.get("dates")
        if not isinstance(dates, list) or not dates:
            raise RequestError("INVALID_EVENT", "设备故障事件的 dates 必须为非空日期数组。")
        parsed = [_parse_date(value, f"dates[{index}]") for index, value in enumerate(dates)]
        if parsed != sorted(parsed) or len(parsed) != len(set(parsed)):
            raise RequestError("INVALID_EVENT", "设备故障日期必须按升序排列且不得重复。")
    if event_type == "absence" and (not isinstance(event.get("employee"), str) or not event["employee"].strip()):
        raise RequestError("INVALID_EVENT", "离岗事件字段 employee 必须为非空员工编号。")
    return event


def validate_request(request: Any, package_root: Path) -> tuple[dict[str, Any], dict[str, Path]]:
    if not isinstance(request, dict):
        raise RequestError("INVALID_REQUEST", "任务请求必须为 JSON 对象。")
    if type(request.get("schema_version")) is not int or request["schema_version"] != REQUEST_SCHEMA_VERSION:
        raise RequestError("INVALID_REQUEST", "schema_version 必须为整数 1。")
    job_id = request.get("job_id")
    if not isinstance(job_id, str) or not JOB_ID_PATTERN.fullmatch(job_id) or ".." in job_id:
        raise RequestError("INVALID_REQUEST", "job_id 只能包含英文字母、数字、点、下划线和连字符，长度为 1 至 64。")
    if not isinstance(request.get("dataset_version"), str) or not request["dataset_version"].strip():
        raise RequestError("INVALID_REQUEST", "dataset_version 必须为非空文本。")
    task = request.get("task")
    if task not in SUPPORTED_TASKS:
        raise RequestError("INVALID_REQUEST", "task 不受支持。", ["支持的任务：" + "、".join(sorted(SUPPORTED_TASKS))])
    raw_files = request.get("input_files")
    if not isinstance(raw_files, dict):
        raise RequestError("INVALID_REQUEST", "input_files 必须为对象。")
    missing = sorted(TASK_FILES[task] - raw_files.keys())
    if missing:
        raise RequestError("INVALID_REQUEST", "任务缺少必需的输入文件引用。", [", ".join(missing)])
    files = {key: _resolve_package_file(package_root, value, key) for key, value in raw_files.items()}
    parameters = request.get("parameters", {})
    if not isinstance(parameters, dict):
        raise RequestError("INVALID_REQUEST", "parameters 必须为对象。")
    normalized = dict(request)
    normalized["parameters"] = parameters
    if task == "reschedule":
        normalized["event"] = _validate_event(request.get("event"))
    if task == "preview_import":
        schema_id = parameters.get("schema_id")
        if schema_id not in IMPORT_SCHEMAS:
            raise RequestError("INVALID_REQUEST", "schema_id 不受支持。", ["请使用系统提供的导入模板类型。"])
    if task == "adjust_schedule":
        adjustment = parameters.get("adjustment")
        if not isinstance(adjustment, dict):
            raise RequestError("INVALID_ADJUSTMENT", "parameters.adjustment 必须为对象。")
        if adjustment.get("type") not in {"assign", "remove"}:
            raise RequestError("INVALID_ADJUSTMENT", "人工调整类型只能为 assign 或 remove。")
        missing_adjustment = [key for key in ("employee_id", "position_code", "date", "shift_code")
                              if not adjustment.get(key)]
        if missing_adjustment:
            raise RequestError("INVALID_ADJUSTMENT", "人工调整缺少必填字段。", missing_adjustment)
        _parse_date(adjustment["date"], "adjustment.date")
        if adjustment["shift_code"] not in {"D12", "N12", "D8", "白班", "夜班", "常白班"}:
            raise RequestError("INVALID_ADJUSTMENT", "shift_code 必须为 D12、N12、D8 或对应班次名称。")
    if task == "validate_schedule":
        profile = parameters.get("demand_profile", "solver")
        if profile not in {"solver", "reference"}:
            raise RequestError("INVALID_REQUEST", "demand_profile 只能为 solver 或 reference。")
        if profile == "reference":
            reference_paths = {
                "position_requirements": "data/reference/input_04_position_requirements.csv",
                "line_run_plan": "data/reference/input_07_line_run_plan.csv",
            }
            for key, default_path in reference_paths.items():
                if key not in files:
                    files[key] = _resolve_package_file(package_root, default_path, key)
    return normalized, files


def _parse_summary(stdout: str) -> Any:
    text = stdout.strip()
    if not text:
        return None
    try:
        value = json.loads(text)
        if isinstance(value, dict):
            return value
    except json.JSONDecodeError:
        pass
    decoder = json.JSONDecoder()
    parsed = []
    for index, char in enumerate(text):
        if char != "{":
            continue
        try:
            value, _ = decoder.raw_decode(text[index:])
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict):
            parsed.append(value)
    if parsed:
        return parsed[-1]
    return {"console_output": text[-4000:]}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _artifact_list(work_dir: Path, package_root: Path) -> list[dict[str, Any]]:
    artifacts = []
    for path in sorted(work_dir.rglob("*")):
        if not path.is_file() or path.name == "event.json":
            continue
        artifacts.append({
            "path": path.resolve().relative_to(package_root.resolve()).as_posix(),
            "sha256": _sha256(path),
            "size_bytes": path.stat().st_size,
        })
    return artifacts


def dispatch_task(request: dict[str, Any], package_root: Path,
                  files: dict[str, Path], work_dir: Path) -> dict[str, Any]:
    """Invoke the existing CLI or independent verifier inside this process."""
    task = request["task"]
    if task == "preview_import":
        preview = preview_import(files["upload"], request["parameters"]["schema_id"])
        return {"return_code": 0, "stdout": json.dumps(preview, ensure_ascii=False), "stderr": ""}
    if task == "validate_schedule":
        result = validate_schedule(
            schedule_path=files["schedule"],
            input_path=files["input"],
            demand_profile=request["parameters"].get("demand_profile", "solver"),
            package_root=package_root,
            position_requirements=files.get("position_requirements"),
            line_run_plan=files.get("line_run_plan"),
            demand_override=files.get("demand_override"),
        )
        return {"return_code": 0 if result["all_pass"] else 4,
                "stdout": json.dumps(result, ensure_ascii=False), "stderr": ""}

    try:
        import sched_solver
    except ImportError as exc:
        raise RequestError("MISSING_DEPENDENCY", "无法加载排班求解依赖。", [str(exc)]) from exc

    args = []
    if task == "info":
        args = ["info", "--input", str(files["input"])]
    elif task in {"solve_demo", "solve_full"}:
        profile = request["parameters"].get("objective_profile", "balanced")
        multiplier = request["parameters"].get("demand_multiplier", 1.0)
        minimum_coverage = request["parameters"].get(
            "minimum_coverage_rate", 0.85 if profile == "cost_control" else None)
        if task == "solve_demo" and profile not in SCHEDULE_OBJECTIVE_PROFILES:
            raise RequestError("INVALID_REQUEST", "objective_profile 不受支持。")
        if task == "solve_demo" and (isinstance(multiplier, bool) or not isinstance(multiplier, (int, float))
                                      or not math.isfinite(multiplier) or not 0.8 <= multiplier <= 1.3):
            raise RequestError("INVALID_REQUEST", "demand_multiplier 必须为 0.8 至 1.3 的有限数值。")
        if task == "solve_demo" and minimum_coverage is not None and (
                isinstance(minimum_coverage, bool) or not isinstance(minimum_coverage, (int, float))
                or not math.isfinite(minimum_coverage) or not 0.5 <= minimum_coverage <= 1.0):
            raise RequestError("INVALID_REQUEST", "minimum_coverage_rate 必须为 0.50 至 1.00 的有限数值。")
        output = work_dir / "schedule.csv"
        args = ["solve-demo" if task == "solve_demo" else "solve-full",
                "--input", str(files["input"]), "--out", str(output),
                "--time-limit", str(request["parameters"].get("time_limit_seconds", 300 if task == "solve_demo" else 600))]
        if task == "solve_demo":
            scenario_name = request["parameters"].get("scenario_name", "PKG-DEMO")
            if not isinstance(scenario_name, str) or len(scenario_name) > 80:
                raise RequestError("INVALID_REQUEST", "scenario_name 必须为 80 个字符以内的文本。")
            args.extend(["--objective-profile", profile, "--demand-multiplier", str(multiplier),
                         "--scenario-name", scenario_name])
            if minimum_coverage is not None:
                args.extend(["--minimum-coverage-rate", str(minimum_coverage)])
    elif task == "verify_baseline":
        args = ["verify-baseline", "--input", str(files["input"]), "--baseline", str(files["baseline"])]
    elif task == "reschedule":
        event_path = work_dir / "event.json"
        event_path.write_text(json.dumps(request["event"], ensure_ascii=False, indent=2), encoding="utf-8")
        args = ["reschedule", "--input", str(files["input"]), "--baseline", str(files["baseline"]),
                "--event", str(event_path), "--out-dir", str(work_dir),
                "--window-days", str(request["parameters"].get("window_days", 7)),
                "--time-limit", str(request["parameters"].get("time_limit_seconds", 45))]
        if "demand_override" in files:
            args.extend(["--demand-override", str(files["demand_override"])])
    elif task == "kpi":
        args = ["kpi", "--baseline", str(files["baseline"]), "--out-dir", str(work_dir)]
    elif task == "adjust_schedule":
        result = adjust_schedule(files["schedule"], files["input"], work_dir / "schedule.csv",
                                 request["parameters"]["adjustment"], files.get("demand_override"))
        return {"return_code": 0 if result["verification"]["all_pass"] else 4,
                "stdout": json.dumps(result, ensure_ascii=False), "stderr": ""}
    else:
        raise RequestError("INVALID_REQUEST", "任务类型没有对应的执行器。")

    stdout, stderr = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            return_code = int(sched_solver.main(args))
    except SystemExit as exc:
        return_code = int(exc.code or 0)
    return {"return_code": return_code, "stdout": stdout.getvalue(), "stderr": stderr.getvalue()}


def _field_error(field: str, code: str, message: str,
                 row: int | None = None) -> dict[str, Any]:
    result = {"field": field, "code": code, "message": message}
    if row is not None:
        result["row"] = row
    return result


def _number_value(value: Any, minimum: float | None = None,
                  maximum: float | None = None, integer: bool = False) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    if minimum is not None and number < minimum:
        return None
    if maximum is not None and number > maximum:
        return None
    if integer and not number.is_integer():
        return None
    return number


def _valid_iso_date(value: Any) -> bool:
    if isinstance(value, (dt.datetime, dt.date)):
        return True
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        return False
    try:
        dt.date.fromisoformat(value.strip())
        return True
    except ValueError:
        return False


def _validate_csv_upload(path: Path, schema_id: str) -> dict[str, Any]:
    schema = IMPORT_SCHEMAS[schema_id]
    errors: list[dict[str, Any]] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
    except UnicodeDecodeError:
        return {"accepted": False, "schema_id": schema_id, "row_count": 0,
                "accepted_rows": 0, "rejected_rows": 0, "columns": [],
                "field_errors": [_field_error("文件编码", "INVALID_ENCODING",
                                               "CSV 文件必须使用 UTF-8 或 UTF-8 BOM 编码。")]}
    if not rows:
        return {"accepted": False, "schema_id": schema_id, "row_count": 0,
                "accepted_rows": 0, "rejected_rows": 0, "columns": [],
                "field_errors": [_field_error("文件", "EMPTY_FILE", "文件没有表头或数据。", 1)]}
    header = [value.strip() for value in rows[0]]
    if len(header) != len(set(header)):
        duplicates = sorted({value for value in header if header.count(value) > 1})
        errors.extend(_field_error(value, "DUPLICATE_HEADER", "表头重复。", 1) for value in duplicates)
    if any(not value for value in header):
        errors.append(_field_error("表头", "EMPTY_HEADER", "表头不能为空。", 1))
    missing = [name for name in schema["required"] if name not in header]
    errors.extend(_field_error(name, "MISSING_FIELD", "缺少必需字段。", 1) for name in missing)
    data_rows = rows[1:]
    row_count = len(data_rows)
    rejected = set()
    for row_number, row in enumerate(data_rows, start=2):
        if len(row) != len(header):
            errors.append(_field_error("列数", "COLUMN_COUNT_MISMATCH", "数据列数与表头列数不一致。", row_number))
            rejected.add(row_number)
        for field in schema["required"]:
            if field in header:
                value = row[header.index(field)].strip() if header.index(field) < len(row) else ""
                if not value:
                    errors.append(_field_error(field, "REQUIRED_VALUE", "必填字段值不能为空。", row_number))
                    rejected.add(row_number)
        if schema_id == "orders" and all(field in header for field in ("required_units", "priority", "due_at")):
            for field in ("required_units", "priority"):
                number = _number_value(row[header.index(field)] if header.index(field) < len(row) else None,
                                       minimum=0, integer=field == "priority")
                if number is None or number <= 0:
                    errors.append(_field_error(field, "INVALID_NUMBER", "该字段必须为正数，priority 必须为整数。", row_number))
                    rejected.add(row_number)
            try:
                dt.datetime.fromisoformat(row[header.index("due_at")].replace("Z", "+00:00"))
            except (ValueError, IndexError):
                errors.append(_field_error("due_at", "INVALID_DATETIME", "该字段必须为 ISO 8601 日期或日期时间。", row_number))
                rejected.add(row_number)
        if schema_id == "line_capacity" and all(field in header for field in ("units_per_hour", "efficiency_pct", "available_hours")):
            for field, minimum, maximum in (("units_per_hour", 0, None), ("efficiency_pct", 0, 100), ("available_hours", 0, 336)):
                cell = row[header.index(field)] if header.index(field) < len(row) else None
                value = _number_value(cell, minimum=minimum, maximum=maximum)
                if value is None or value <= minimum:
                    errors.append(_field_error(field, "INVALID_RANGE", "数值超出允许范围。", row_number))
                    rejected.add(row_number)
            shift = row[header.index("shift_code")].strip() if header.index("shift_code") < len(row) else ""
            if shift not in {"D12", "N12", "D8", "白班", "夜班", "常白班"}:
                errors.append(_field_error("shift_code", "INVALID_ENUM", "班次代码不受支持。", row_number))
                rejected.add(row_number)
        if schema_id in {"line_run_plan", "attendance_snapshot", "baseline_schedule"}:
            date_field = "日期" if schema_id in {"line_run_plan", "baseline_schedule"} else "schedule_date"
            if date_field in header:
                raw_date = row[header.index(date_field)].strip() if header.index(date_field) < len(row) else ""
                if not _valid_iso_date(raw_date):
                    errors.append(_field_error(date_field, "INVALID_DATE", "该字段必须为有效的 YYYY-MM-DD 日期。", row_number))
                    rejected.add(row_number)
        if schema_id == "line_run_plan":
            shift = row[header.index("班次")].strip() if "班次" in header and header.index("班次") < len(row) else ""
            status = row[header.index("生产运行状态")].strip() if "生产运行状态" in header and header.index("生产运行状态") < len(row) else ""
            if shift not in {"白班", "夜班", "常白班"}:
                errors.append(_field_error("班次", "INVALID_ENUM", "班次必须为白班、夜班或常白班。", row_number))
                rejected.add(row_number)
            if status not in {"正常生产运行", "换型", "大修", "例行维护", "停线", "停产"}:
                errors.append(_field_error("生产运行状态", "INVALID_ENUM", "生产运行状态不受支持。", row_number))
                rejected.add(row_number)
        if schema_id == "attendance_snapshot":
            shift = row[header.index("shift_code")].strip() if "shift_code" in header and header.index("shift_code") < len(row) else ""
            status = row[header.index("status")].strip() if "status" in header and header.index("status") < len(row) else ""
            if shift not in {"D12", "N12", "D8", "白班", "夜班", "常白班"}:
                errors.append(_field_error("shift_code", "INVALID_ENUM", "班次代码不受支持。", row_number))
                rejected.add(row_number)
            if status not in {"present", "absent", "leave", "到岗", "缺勤", "请假"}:
                errors.append(_field_error("status", "INVALID_ENUM", "考勤状态不受支持。", row_number))
                rejected.add(row_number)
        if schema_id == "position_requirements":
            for field, maximum in (("岗位技能要求等级", 10), ("白班定编", 500), ("夜班定编", 500)):
                if field not in header:
                    continue
                cell = row[header.index(field)] if header.index(field) < len(row) else None
                if _number_value(cell, minimum=0, maximum=maximum, integer=True) is None:
                    errors.append(_field_error(field, "INVALID_RANGE", "该字段必须为允许范围内的非负整数。", row_number))
                    rejected.add(row_number)
        if schema_id == "baseline_schedule":
            code = row[header.index("班次代码")].strip() if "班次代码" in header and header.index("班次代码") < len(row) else ""
            if code not in {"D12", "N12", "D8"}:
                errors.append(_field_error("班次代码", "INVALID_ENUM", "班次代码必须为 D12、N12 或 D8。", row_number))
                rejected.add(row_number)
            for field in ("计薪工时(h)", "标准工时(h)", "当日加班工时(h)"):
                if field in header:
                    cell = row[header.index(field)] if header.index(field) < len(row) else None
                    if _number_value(cell, minimum=0, maximum=24) is None:
                        errors.append(_field_error(field, "INVALID_RANGE", "工时必须为 0 至 24 h 的有限数值。", row_number))
                        rejected.add(row_number)
    unique_fields = CSV_UNIQUE_FIELDS.get(schema_id, ())
    if unique_fields and all(field in header for field in unique_fields):
        seen = set()
        for row_number, row in enumerate(data_rows, start=2):
            key = tuple(row[header.index(field)].strip() if header.index(field) < len(row) else ""
                        for field in unique_fields)
            if any(not value for value in key):
                continue
            if key in seen:
                errors.append(_field_error(",".join(unique_fields), "DUPLICATE_RECORD", "数据记录键重复。", row_number))
                rejected.add(row_number)
            seen.add(key)
    header_rejected = bool(missing or any(error["code"] in {"DUPLICATE_HEADER", "EMPTY_HEADER"} for error in errors))
    if header_rejected:
        rejected.update(range(2, row_count + 2))
    return {
        "accepted": not errors and row_count > 0,
        "schema_id": schema_id,
        "row_count": row_count,
        "accepted_rows": max(0, row_count - len(rejected)),
        "rejected_rows": len(rejected),
        "columns": header,
        "field_errors": errors[:200],
    }


def preview_import(path: Path, schema_id: str) -> dict[str, Any]:
    """Preview supported files; validation errors are data, not task crashes."""
    schema = IMPORT_SCHEMAS.get(schema_id)
    if schema is None:
        raise RequestError("INVALID_REQUEST", "schema_id 不受支持。")
    suffix = path.suffix.lower()
    if suffix != schema["extension"]:
        return {"accepted": False, "schema_id": schema_id, "row_count": 0,
                "accepted_rows": 0, "rejected_rows": 0, "columns": [],
                "field_errors": [_field_error("文件类型", "UNSUPPORTED_EXTENSION",
                                               f"{schema_id} 模板只接受 {schema['extension']} 文件。")]}
    if suffix == ".csv":
        return _validate_csv_upload(path, schema_id)
    if schema_id == "event_json":
        try:
            event = json.loads(path.read_text(encoding="utf-8"))
            _validate_event(event)
            return {"accepted": True, "schema_id": schema_id, "row_count": 1,
                    "accepted_rows": 1, "rejected_rows": 0, "columns": list(event), "field_errors": []}
        except (OSError, json.JSONDecodeError, RequestError) as exc:
            return {"accepted": False, "schema_id": schema_id, "row_count": 0,
                    "accepted_rows": 0, "rejected_rows": 0, "columns": [],
                    "field_errors": [_field_error("JSON", "INVALID_JSON", str(exc))]}
    if schema_id == "raw_dataset":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RequestError("MISSING_DEPENDENCY", "读取 XLSX 文件需要 openpyxl。", [str(exc)]) from exc
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            return {"accepted": False, "schema_id": schema_id, "row_count": 0,
                    "accepted_rows": 0, "rejected_rows": 0, "columns": [],
                    "field_errors": [_field_error("文件", "INVALID_XLSX", "文件不是可读取的 XLSX 工作簿。")]}
        errors = []
        total_rows = 0
        accepted_rows = 0
        rejected_rows = 0
        sheets = {}
        for sheet_name, required in RAW_WORKBOOK_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                errors.append(_field_error(sheet_name, "MISSING_SHEET", "缺少必需工作表。"))
                continue
            sheet = workbook[sheet_name]
            headers = [
                str(value).strip() if value is not None else ""
                for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            ]
            missing = [name for name in required if name not in headers]
            errors.extend(_field_error(f"{sheet_name}.{name}", "MISSING_FIELD", "工作表缺少必需字段。", 1)
                          for name in missing)
            duplicate_headers = sorted({value for value in headers if value and headers.count(value) > 1})
            errors.extend(_field_error(f"{sheet_name}.{name}", "DUPLICATE_HEADER", "工作表表头重复。", 1)
                          for name in duplicate_headers)
            row_count = 0
            sheet_accepted = 0
            sheet_rejected = 0
            header_index = {name: index for index, name in enumerate(headers)}
            numeric_fields = {
                "设备产线与工位台账": (("岗位技能要求等级", 0, 10, True), ("白班定编", 0, 500, True), ("夜班定编", 0, 500, True)),
                "人员技能矩阵": (("技能等级", 0, 10, True), ("累计在岗时长", 0, None, False)),
                "排班训练主数据集": (("换型等级", 0, 10, True),),
            }.get(sheet_name, ())
            for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                row_count += 1
                row_errors = []
                for field in required:
                    index = header_index.get(field)
                    value = values[index] if index is not None and index < len(values) else None
                    if value is None or (isinstance(value, str) and not value.strip()):
                        row_errors.append(_field_error(f"{sheet_name}.{field}", "REQUIRED_VALUE", "必填字段值不能为空。", row_number))
                for field, minimum, maximum, integer in numeric_fields:
                    index = header_index.get(field)
                    value = values[index] if index is not None and index < len(values) else None
                    checked = _number_value(value, minimum=minimum, maximum=maximum, integer=integer)
                    if checked is None:
                        row_errors.append(_field_error(f"{sheet_name}.{field}", "INVALID_RANGE", "该字段必须为允许范围内的非负整数。", row_number))
                date_fields = ("日期",) if sheet_name == "排班训练主数据集" else ()
                for field in date_fields:
                    index = header_index.get(field)
                    value = values[index] if index is not None and index < len(values) else None
                    if not _valid_iso_date(value):
                        row_errors.append(_field_error(f"{sheet_name}.{field}", "INVALID_DATE", "该字段必须为有效日期。", row_number))
                if row_errors:
                    sheet_rejected += 1
                    errors.extend(row_errors)
                else:
                    sheet_accepted += 1
            if not row_count:
                errors.append(_field_error(sheet_name, "EMPTY_SHEET", "工作表至少需要一条有效数据记录。"))
            total_rows += row_count
            accepted_rows += sheet_accepted
            rejected_rows += sheet_rejected
            sheets[sheet_name] = {"row_count": row_count, "accepted_rows": sheet_accepted,
                                  "rejected_rows": sheet_rejected, "columns": headers}
        workbook.close()
        return {"accepted": not errors and total_rows > 0, "schema_id": schema_id,
                "row_count": total_rows, "accepted_rows": accepted_rows,
                "rejected_rows": rejected_rows, "columns": [],
                "sheets": sheets, "field_errors": errors[:200]}
    return {"accepted": False, "schema_id": schema_id, "row_count": 0,
            "accepted_rows": 0, "rejected_rows": 0, "columns": [],
            "field_errors": [_field_error("文件", "UNSUPPORTED_SCHEMA", "文件格式或模板类型不受支持。")]}


def _load_demand_override(path: Path, data: dict[str, Any]) -> dict[tuple[str, str, str], float]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload["effective_demand"]
        demand = {}
        valid_positions = set(data["POS"])
        valid_dates = set(data["dates"])
        valid_shifts = {"白班", "夜班", "常白班"}
        for row in rows:
            key = (row["position_code"], row["date"], row["shift"])
            value = _number_value(row["demand"], minimum=0)
            if key[0] not in valid_positions or key[1] not in valid_dates or key[2] not in valid_shifts:
                raise ValueError("需求单元超出当前数据集范围")
            if value is None or key in demand:
                raise ValueError("需求数量无效或需求单元重复")
            demand[key] = value
        if not demand:
            raise ValueError("有效需求列表为空")
        return demand
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise RequestError("INVALID_DEMAND_OVERRIDE", "排班版本的有效需求数据无效。", [str(exc)]) from exc


def adjust_schedule(schedule_path: Path, input_path: Path, output_path: Path,
                    adjustment: dict[str, Any], demand_override_path: Path | None = None) -> dict[str, Any]:
    """Apply one assignment change and run the independent hard-constraint verifier."""
    src_dir = Path(__file__).resolve().parent
    if str(src_dir) not in sys.path:
        sys.path.insert(0, str(src_dir))
    import scheduling_model_builder as builder
    import sched_solver
    from verify_core import NAME2CODE, SHIFT_META, kpi_components, parse_schedule_csv, verify_seven

    data = builder.load_data(str(input_path))
    if demand_override_path is not None:
        data["demand"] = _load_demand_override(demand_override_path, data)
    employee = str(adjustment["employee_id"]).strip()
    position = str(adjustment["position_code"]).strip()
    schedule_date = str(adjustment["date"])
    shift_code = NAME2CODE.get(adjustment["shift_code"], adjustment["shift_code"])
    if employee not in data["emp_list"]:
        raise RequestError("INVALID_ADJUSTMENT", "员工编号不存在于所选数据集。", [employee])
    if position not in data["POS"]:
        raise RequestError("INVALID_ADJUSTMENT", "岗位编码不存在于所选数据集。", [position])
    if schedule_date not in data["dates"]:
        raise RequestError("INVALID_ADJUSTMENT", "排班日期超出所选数据集范围。", [schedule_date])
    if shift_code not in SHIFT_META:
        raise RequestError("INVALID_ADJUSTMENT", "班次代码无效。", [shift_code])

    assign, gaps, fvals = parse_schedule_csv(str(schedule_path))
    key = (position, schedule_date, SHIFT_META[shift_code]["name"])
    target = (employee, position, schedule_date, shift_code)
    if adjustment["type"] == "assign":
        if target in assign:
            raise RequestError("INVALID_ADJUSTMENT", "该员工已存在相同岗位、日期和班次指派。")
        if gaps.get(key, 0) > 0:
            gaps[key] -= 1
        assign.append(target)
        action = "新增指派"
    else:
        if target not in assign:
            raise RequestError("INVALID_ADJUSTMENT", "未找到可取消的对应指派。")
        assign.remove(target)
        gaps[key] = gaps.get(key, 0) + 1
        action = "取消指派"

    sched_solver.write_schedule_csv(str(output_path), assign, gaps, fvals, data, "MANUAL-EDIT")
    all_pass, violations, samples = verify_seven(assign, gaps, fvals, data)
    kpi = kpi_components(assign, gaps, fvals, data)
    effective_demand_path = output_path.with_name("effective_demand.json")
    effective_demand_path.write_text(json.dumps({
        "effective_demand": [
            {"position_code": key[0], "date": key[1], "shift": key[2], "demand": value}
            for key, value in sorted(data["demand"].items())
        ]
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "action": action,
        "change": {"employee_id": employee, "position_code": position,
                   "date": schedule_date, "shift_code": shift_code},
        "verification": {"all_pass": all_pass, "violations": violations, "samples": samples[:10]},
        "kpi": kpi,
        "output": output_path.as_posix(),
        "output_sha256": sched_solver.sha256_file(output_path),
        "effective_demand": effective_demand_path.as_posix(),
        "note": "调整后已运行独立硬约束核验；核验失败的方案不可审批或锁定。",
    }


def validate_schedule(schedule_path: Path, input_path: Path, demand_profile: str,
                      package_root: Path | None = None,
                      position_requirements: Path | None = None,
                      line_run_plan: Path | None = None,
                      demand_override: Path | None = None) -> dict[str, Any]:
    """Run independent hard-constraint verification, without KPI baseline checks."""
    src_dir = Path(__file__).resolve().parent
    if str(src_dir) not in sys.path:
        sys.path.insert(0, str(src_dir))
    import scheduling_model_builder as builder
    from verify_core import parse_schedule_csv, verify_seven

    data = builder.load_data(str(input_path))
    if demand_override is not None:
        if demand_profile != "solver":
            raise RequestError("INVALID_REQUEST", "基于事件的方案版本只能使用其保存的有效需求口径核验。")
        demand = _load_demand_override(demand_override, data)
    else:
        demand = None
    if demand_profile == "solver":
        demand = demand or data["demand"]
    elif demand_profile == "reference":
        if position_requirements is None or line_run_plan is None:
            raise RequestError("INVALID_REQUEST", "reference 需求口径必须提供岗位定编和产线运行计划文件。")
        reference_data = builder.load_data_from_csv(
            str(input_path), str(position_requirements), str(line_run_plan))
        demand = reference_data["demand"]
    else:
        raise RequestError("INVALID_REQUEST", "需求口径必须为 solver 或 reference。")

    assign, gaps, fvals = parse_schedule_csv(str(schedule_path))
    all_pass, violations, samples = verify_seven(assign, gaps, fvals, data, demand)
    return {
        "demand_profile": demand_profile,
        "all_pass": all_pass,
        "violations": violations,
        "samples": samples[:10],
        "violation_definitions": HARD_CONSTRAINT_LABELS,
        "assignment_count": len(assign),
        "demand_cells": len(demand),
        "note": "此结果仅报告独立硬约束核验，不比较 KPI 基准登记值。",
    }


def execute(request: dict[str, Any], package_root: Path) -> dict[str, Any]:
    """Validate a job request, dispatch it, and produce one stable response."""
    start = time.perf_counter()
    root = Path(package_root)
    job_id = _safe_job_id(request.get("job_id") if isinstance(request, dict) else None)
    try:
        normalized, files = validate_request(request, root)
        job_id = normalized["job_id"]
    except RequestError as exc:
        return _error_response(job_id, exc.code, str(exc), exc.details)
    except (OSError, RuntimeError) as exc:
        return _error_response(job_id, "PACKAGE_ROOT_ERROR", "无法读取工程包目录。", [str(exc)])

    # Validate time limits and event references before creating a task directory.
    params = normalized["parameters"]
    time_limit = params.get("time_limit_seconds")
    if time_limit is not None:
        if isinstance(time_limit, bool) or not isinstance(time_limit, (int, float)) or time_limit <= 0:
            return _error_response(job_id, "INVALID_REQUEST", "time_limit_seconds 必须为正数。")
    if normalized["task"] == "reschedule":
        window_days = params.get("window_days", 7)
        if isinstance(window_days, bool) or not isinstance(window_days, int) or not 1 <= window_days <= 62:
            return _error_response(job_id, "INVALID_REQUEST", "window_days 必须为 1 至 62 的整数。")
        try:
            _validate_event_references(normalized["event"], files["input"])
        except RequestError as exc:
            return _error_response(job_id, exc.code, str(exc), exc.details)
        except ImportError as exc:
            return _error_response(job_id, "MISSING_DEPENDENCY", "无法校验事件所引用的数据。", [str(exc)])

    root_resolved = root.resolve()
    job_dir = root_resolved / "out" / "jobs" / job_id
    work_dir = job_dir / "artifacts"
    try:
        job_dir_resolved = job_dir.resolve(strict=False)
        job_dir_resolved.relative_to(root_resolved)
        work_dir.mkdir(parents=True, exist_ok=False)
    except ValueError:
        return _error_response(job_id, "OUTPUT_PATH_OUTSIDE_PACKAGE", "任务输出目录超出工程包目录。")
    except FileExistsError:
        return _error_response(job_id, "JOB_OUTPUT_EXISTS", "该任务编号已有输出目录，请使用新的 job_id。")
    except OSError as exc:
        return _error_response(job_id, "OUTPUT_DIRECTORY_ERROR", "无法创建任务输出目录。", [str(exc)])
    try:
        result = dispatch_task(normalized, root_resolved, files, work_dir)
    except RequestError as exc:
        result = {"return_code": 2, "stdout": "", "stderr": str(exc), "error": exc}
    except Exception as exc:  # Convert failures into the Rust process contract.
        result = {"return_code": 1, "stdout": "", "stderr": f"{type(exc).__name__}: {exc}"}

    return_code = int(result.get("return_code", 1))
    stdout = str(result.get("stdout", ""))
    stderr = str(result.get("stderr", ""))
    (work_dir / "stdout.log").write_text(stdout, encoding="utf-8")
    (work_dir / "stderr.log").write_text(stderr, encoding="utf-8")
    summary = _parse_summary(stdout)
    verification = summary.get("verification") if isinstance(summary, dict) else None
    if verification is None and isinstance(summary, dict) and "all_pass" in summary:
        verification = summary
    error = None
    if return_code != 0:
        if isinstance(result.get("error"), RequestError):
            err = result["error"]
            error = {"code": err.code, "message": str(err), "details": err.details}
        else:
            error = {
                "code": "TASK_FAILED",
                "message": "任务执行失败。",
                "details": [stderr.strip()[-1000:]] if stderr.strip() else [],
            }
    return _response(
        job_id, "succeeded" if return_code == 0 else "failed", return_code,
        summary=summary, verification=verification,
        timings={"wall_seconds": round(time.perf_counter() - start, 3)},
        artifacts=_artifact_list(work_dir, root_resolved), error=error,
    )


def _validate_event_references(event: dict[str, Any], input_path: Path) -> None:
    """Reject dates and entity identifiers that cannot occur in the selected dataset."""
    src_dir = Path(__file__).resolve().parent
    if str(src_dir) not in sys.path:
        sys.path.insert(0, str(src_dir))
    import scheduling_model_builder as builder

    data = builder.load_data(str(input_path))
    valid_dates = set(data["dates"])
    dates = event["dates"] if event["type"] == "equipment_fail" else [event["date"]]
    outside = [value for value in dates if value not in valid_dates]
    if outside:
        raise RequestError("INVALID_EVENT", "事件日期超出数据集排班周期。", outside)
    if event["type"] in {"insert", "rampup", "changeover", "equipment_fail"}:
        if event["line"] not in data["lines"]:
            raise RequestError("INVALID_EVENT", "事件产线编码不存在于所选数据集。", [event["line"]])
    if event["type"] == "absence" and event["employee"] not in data["emp_list"]:
        raise RequestError("INVALID_EVENT", "离岗员工编号不存在于所选数据集。", [event["employee"]])


def preflight_event(event: Any, input_path: Path) -> dict[str, Any]:
    """Validate event structure and dataset references without creating a job."""
    normalized = _validate_event(event)
    try:
        _validate_event_references(normalized, input_path)
    except ImportError as exc:
        raise RequestError("MISSING_DEPENDENCY", "无法校验事件所引用的数据。", [str(exc)]) from exc
    return {"valid": True, "event_id": normalized.get("event_id", "EVT")}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Rust 调度的本机排班任务适配器")
    parser.add_argument("--request", help="包内任务 JSON 路径")
    parser.add_argument("--preflight-event", action="store_true", help="仅执行事件结构与数据引用预检")
    parser.add_argument("--input", help="包内数据集输入文件")
    parser.add_argument("--event-json", help="事件 JSON 文本")
    parser.add_argument("--package-root", required=True, help="工程包根目录")
    args = parser.parse_args(argv)
    try:
        if args.preflight_event:
            if not args.input or not args.event_json:
                raise RequestError("INVALID_REQUEST", "事件预检必须提供 input 与 event-json。")
            input_path = _resolve_package_file(Path(args.package_root), args.input, "input")
            event = json.loads(args.event_json)
            response = preflight_event(event, input_path)
        else:
            if not args.request:
                raise RequestError("INVALID_REQUEST", "必须提供任务 request 文件。")
            request_path = _resolve_package_file(Path(args.package_root), args.request, "request")
            request = json.loads(request_path.read_text(encoding="utf-8"))
            response = execute(request, Path(args.package_root))
    except RequestError as exc:
        response = {"valid": False, "error": {"code": exc.code, "message": str(exc), "details": exc.details}} \
            if args.preflight_event else _error_response("invalid-request", exc.code, str(exc), exc.details)
    except (OSError, json.JSONDecodeError) as exc:
        response = {"valid": False, "error": {"code": "INVALID_REQUEST", "message": "无法读取事件预检输入。", "details": [str(exc)]}} \
            if args.preflight_event else _error_response("invalid-request", "INVALID_REQUEST", "无法读取任务请求 JSON。", [str(exc)])
    print(json.dumps(response, ensure_ascii=False, separators=(",", ":")))
    return (0 if response.get("valid") else 2) if args.preflight_event else int(response["exit_code"])


if __name__ == "__main__":
    sys.exit(main())
