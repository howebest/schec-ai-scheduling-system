#!/usr/bin/env python3
"""Build a small, synthetic workbook for the standalone Streamlit demo."""
from __future__ import annotations

import argparse
import csv
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
POSITIONS_CSV = ROOT / "data" / "reference" / "input_04_position_requirements.csv"
EMPLOYEES_CSV = ROOT / "data" / "demo" / "expanded" / "employees.csv"
DEFAULT_OUTPUT = ROOT / "data" / "demo" / "streamlit_demo_dataset.xlsx"
LINE_CODE = "SX-PET01"
FACTORY_CODE = "SX"
DEMO_START = date(2026, 10, 12)

LEDGER_COLUMNS = [
    "工厂代码", "产线编码", "岗位编码", "岗位名称", "岗位类别",
    "岗位技能要求等级", "白班定编", "夜班定编", "轮班模式",
    "岗位证书要求",
]
SKILL_COLUMNS = [
    "员工工号", "姓名", "工厂代码", "技能岗位", "技能等级",
    "证书有效期", "可胜任班次",
]
TRAINING_COLUMNS = [
    "工厂代码", "产线编码", "日期", "班次", "生产运行状态",
    "换型标记", "换型等级", "岗位编码", "持证要求",
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def create_workbook(output: Path) -> None:
    positions = [
        row for row in read_csv(POSITIONS_CSV)
        if row["产线编码"] == LINE_CODE and row["工厂代码"] == FACTORY_CODE
    ]
    employees = [
        row for row in read_csv(EMPLOYEES_CSV)
        if row["factory_code"] == FACTORY_CODE and row["employment_status"] == "在岗"
    ]
    employees = sorted(employees, key=lambda row: row["employee_id"])[:42]
    if not positions or len(employees) < len(positions) * 3:
        raise RuntimeError("合成演示工作簿所需岗位或人员数据不足。")

    workbook = Workbook()
    ledger_sheet = workbook.active
    ledger_sheet.title = "设备产线与工位台账"
    ledger_sheet.append(LEDGER_COLUMNS)
    for row in positions:
        ledger_sheet.append([
            row["工厂代码"], row["产线编码"], row["岗位编码"], row["岗位名称"],
            row["岗位类别"], int(row["岗位技能要求等级"]), int(row["白班定编"]),
            int(row["夜班定编"]), row["轮班模式"], row["岗位证书要求"],
        ])

    skill_sheet = workbook.create_sheet("人员技能矩阵")
    skill_sheet.append(SKILL_COLUMNS)
    ordered_positions = sorted(positions, key=lambda row: row["岗位编码"])
    for employee_index, employee in enumerate(employees):
        # 每人配置三项合成技能，保证各岗位有可用候选人员，同时形成多技能组合。
        position_indices = {
            employee_index % len(ordered_positions),
            (employee_index + 2) % len(ordered_positions),
            (employee_index + 4) % len(ordered_positions),
        }
        night_eligible = employee_index % 5 != 0
        for position_index in sorted(position_indices):
            position = ordered_positions[position_index]
            required_level = int(position["岗位技能要求等级"])
            skill_level = max(required_level, 3 + (employee_index % 2))
            skill_sheet.append([
                employee["employee_id"], employee["name"], FACTORY_CODE,
                position["岗位名称"], skill_level, "",
                "白班、夜班" if night_eligible else "白班",
            ])

    plan_sheet = workbook.create_sheet("排班训练主数据集")
    plan_sheet.append(TRAINING_COLUMNS)
    for day_offset in range(7):
        schedule_date = (DEMO_START + timedelta(days=day_offset)).isoformat()
        for shift in ("白班", "夜班"):
            for position in positions:
                plan_sheet.append([
                    FACTORY_CODE, LINE_CODE, schedule_date, shift, "正常生产运行",
                    0, "", position["岗位编码"], position["岗位证书要求"],
                ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    create_workbook(args.out)
    print(f"合成演示工作簿已生成：{args.out}")


if __name__ == "__main__":
    main()
